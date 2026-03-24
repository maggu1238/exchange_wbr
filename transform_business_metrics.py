#!/usr/bin/env python3
"""
Business Metrics: Transform raw xlsx data into a WBR-style report.

Reads 7 file types (dc, comp100, comp75, depth, attach, reseller, sweetner)
across 3 timeframes (weekly, monthly, pytd) and produces a single CSV report.

Can be run standalone:
    python3 transform_business_metrics.py

Or called by the dispatcher (run.py) via the run() entry point.
"""

import csv
import re
import argparse
import yaml
import openpyxl
from pathlib import Path
from datetime import date

METRIC_NAME = "business_metrics"

MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
NUM_TO_MONTH = {i: m for i, m in enumerate(MONTH_ABBR, 1)}

MAX_WEEKLY_PERIODS = 6
MAX_MONTHLY_PERIODS = 6

# ═══════════════════════════════════════════════════════════════════
#  TIME-FRAME PARSERS
# ═══════════════════════════════════════════════════════════════════

def _parse_weekly(raw):
    """'2026-9' -> sort_key=(2026,9), label='WK-09'"""
    parts = str(raw).split("-")
    year, week = int(parts[0]), int(parts[1])
    return (year, week), f"WK-{week:02d}"


def _parse_monthly(raw):
    """Handles both '2026-3' and '3-2026' -> sort_key=(2026,3), label=\"Mar'26\" """
    parts = str(raw).split("-")
    a, b = int(parts[0]), int(parts[1])
    if a > 12:
        year, month = a, b
    else:
        month, year = a, b
    label = f"{NUM_TO_MONTH[month]}'{year % 100:02d}"
    return (year, month), label


def _sort_labels(period_map):
    """Given {label: sort_key}, return labels sorted descending."""
    return sorted(period_map.keys(), key=lambda l: period_map[l], reverse=True)


# ═══════════════════════════════════════════════════════════════════
#  LOADERS  — one per file format
#  Each returns: {display_device: {period_label: value_or_dict}}
# ═══════════════════════════════════════════════════════════════════

def _build_alias_map(aliases_cfg, file_type):
    """Build {input_name_lower: display_name} from config."""
    raw = aliases_cfg.get(file_type, {})
    return {str(k).lower().strip(): v for k, v in raw.items()}


def _load_format_a(filepath, parse_fn, alias_map):
    """Format A: Group | subcat | period1 | period2 | ...
    Used by: dc, attach, reseller.
    Returns {display_device: {label: value}} and {label: sort_key}.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_row = None
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() if c else "" for c in row]
        if "subcat" in cells or "group" in cells:
            header_row = i
            break
    if header_row is None:
        return {}, {}

    headers = list(rows[header_row])
    period_cols = []
    period_map = {}
    for ci in range(2, len(headers)):
        raw = headers[ci]
        if raw is None:
            continue
        raw_s = str(raw).strip()
        if raw_s.lower() in ("ytd", "total"):
            period_cols.append((ci, "__total__"))
            continue
        try:
            sk, label = parse_fn(raw_s)
            period_cols.append((ci, label))
            period_map[label] = sk
        except (ValueError, KeyError, IndexError):
            continue

    skip = {"hctp subtotal", "ohl subtotal", "total", ""}
    data = {}
    for row in rows[header_row + 1:]:
        subcat = str(row[1]).strip() if row[1] is not None else ""
        if subcat.lower() in skip or subcat == "":
            continue
        display = alias_map.get(subcat.lower(), subcat)
        vals = {}
        for ci, label in period_cols:
            v = row[ci]
            if v is not None:
                try:
                    vals[label] = float(v)
                except (ValueError, TypeError):
                    pass
        data[display] = vals

    return data, period_map


def _load_format_b(filepath, parse_fn, alias_map, sub_cols):
    """Format B: Category | (sub_col1, sub_col2, ...) × period
    Used by: comp100 (Ovrl/Base/Swtnr), comp75 (Ovrl/Base).
    Returns {display_device: {label: {sub_col: value}}} and {label: sort_key}.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header0 = list(rows[0])
    n_sub = len(sub_cols)

    period_cols = []
    period_map = {}
    for ci, val in enumerate(header0):
        if val is None or ci == 0:
            continue
        raw_s = str(val).strip()
        if raw_s.lower() == "total":
            period_cols.append((ci, "__total__"))
            continue
        try:
            sk, label = parse_fn(raw_s)
            period_cols.append((ci, label))
            period_map[label] = sk
        except (ValueError, KeyError, IndexError):
            continue

    data = {}
    for row in rows[2:]:
        cat = str(row[0]).strip() if row[0] is not None else ""
        if cat.lower() == "total" or cat == "":
            continue
        display = alias_map.get(cat.lower(), cat)
        vals = {}
        for ci, label in period_cols:
            sub_vals = {}
            for si, sc in enumerate(sub_cols):
                idx = ci + si
                if idx < len(row) and row[idx] is not None:
                    try:
                        sub_vals[sc] = float(row[idx])
                    except (ValueError, TypeError):
                        pass
            if sub_vals:
                vals[label] = sub_vals
        data[display] = vals

    return data, period_map


def _load_format_d(filepath, parse_fn, alias_map):
    """Format D: Category | period1 | period2 | ...
    Used by: depth. Single value per period.
    Returns {display_device: {label: value}} and {label: sort_key}.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = list(rows[0])
    period_cols = []
    period_map = {}
    for ci in range(1, len(header)):
        raw = header[ci]
        if raw is None:
            continue
        raw_s = str(raw).strip()
        if raw_s.lower() == "total":
            period_cols.append((ci, "__total__"))
            continue
        try:
            sk, label = parse_fn(raw_s)
            period_cols.append((ci, label))
            period_map[label] = sk
        except (ValueError, KeyError, IndexError):
            continue

    data = {}
    for row in rows[1:]:
        cat = str(row[0]).strip() if row[0] is not None else ""
        if cat.lower() == "total" or cat == "":
            continue
        display = alias_map.get(cat.lower(), cat)
        vals = {}
        for ci, label in period_cols:
            v = row[ci]
            if v is not None:
                try:
                    vals[label] = float(v)
                except (ValueError, TypeError):
                    pass
        data[display] = vals

    return data, period_map


def _load_format_e(filepath, parse_fn, alias_map):
    """Format E: Sweetner coverage — pivot with (sponsor, vm_session) per period.
    Row 4 has period labels, Row 5 has sub-column names.
    Returns {display_device: {label: {sponsor: v, vm_session: v}}} and {label: sort_key}.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    period_row = None
    header_row = None
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() if c else "" for c in row]
        if "rows" in cells and "sponsor" in cells:
            header_row = i
            period_row = i - 1
            break
    if header_row is None:
        return {}, {}

    period_header = list(rows[period_row])
    sub_header = list(rows[header_row])

    period_cols = []
    period_map = {}
    ci = 1
    while ci < len(period_header):
        raw = period_header[ci]
        if raw is not None:
            raw_s = str(raw).strip()
            if raw_s.lower() == "total":
                period_cols.append((ci, "__total__"))
                ci += 2
                continue
            try:
                sk, label = parse_fn(raw_s)
                period_cols.append((ci, label))
                period_map[label] = sk
            except (ValueError, KeyError, IndexError):
                ci += 1
                continue
            ci += 2
        else:
            ci += 1

    skip = {"total", "no subcat", ""}
    data = {}
    for row in rows[header_row + 1:]:
        cat = str(row[0]).strip() if row[0] is not None else ""
        if cat.lower() in skip or cat == "":
            continue
        display = alias_map.get(cat.lower(), cat)
        vals = {}
        for col_start, label in period_cols:
            sub_vals = {}
            for offset, sc in enumerate(["sponsor", "vm_session"]):
                idx = col_start + offset
                if idx < len(row) and row[idx] is not None:
                    try:
                        sub_vals[sc] = float(row[idx])
                    except (ValueError, TypeError):
                        pass
            if sub_vals:
                vals[label] = sub_vals
        data[display] = vals

    return data, period_map


# ═══════════════════════════════════════════════════════════════════
#  PREVIOUS REPORT READER  (for OP2 carry-forward)
# ═══════════════════════════════════════════════════════════════════

def _header_sections(header, start=4):
    """Split header columns (from `start`) into sections delimited by empty cells.
    Returns list of (col_start, length) tuples for each non-empty section.
    """
    sections = []
    i = start
    while i < len(header):
        name = str(header[i]).strip() if header[i] else ""
        if name:
            s = i
            while i < len(header):
                n = str(header[i]).strip() if header[i] else ""
                if not n:
                    break
                i += 1
            sections.append((s, i - s))
        else:
            i += 1
    return sections


def _repair_csv_row(row, expected_cols):
    """Repair a CSV row where numbers with thousands separators got split.
    e.g. "1,083" parsed as ["1", "083"] is merged back to ["1083"].
    """
    if len(row) <= expected_cols:
        return row
    repaired = []
    i = 0
    while i < len(row):
        val = row[i].strip()
        if (i + 1 < len(row) and
                val.lstrip("-").isdigit() and
                re.match(r"^\d{3}$", row[i + 1].strip())):
            repaired.append(val + row[i + 1].strip())
            i += 2
        else:
            repaired.append(row[i])
            i += 1
    return repaired


def _load_previous_report(filepath):
    """Read previous CSV output. Auto-detects header row and column layout.
    Returns:
    - op2_values: {(category, metric): op2_value}
    - ytd_op2_values: {(category, metric): ytd_op2_value}
    - hardcoded_rows: {(category, metric): full_row_list}
    """
    op2 = {}
    ytd_op2 = {}
    hardcoded = {}

    if not Path(filepath).is_file():
        return op2, ytd_op2, hardcoded

    with open(filepath, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if len(rows) < 3:
        return op2, ytd_op2, hardcoded

    header_idx = None
    cat_col = None
    met_col = None
    for ri, row in enumerate(rows[:10]):
        cells = [str(c).strip() for c in row]
        for ci, c in enumerate(cells):
            if c == "Category":
                cat_col = ci
            if c in ("Metrics", "Metric"):
                met_col = ci
        if cat_col is not None and met_col is not None:
            header_idx = ri
            break

    if header_idx is None:
        return op2, ytd_op2, hardcoded

    header = rows[header_idx]
    expected_cols = len(header)
    op2_col = None
    ytd_op2_col = None
    for ci, h in enumerate(header):
        h_s = str(h).strip()
        if "OP2" in h_s and "YTD" not in h_s and "v/s" not in h_s and "July" not in h_s:
            op2_col = ci
        if "YTD OP2" in h_s:
            ytd_op2_col = ci

    prev_sections = _header_sections(header)

    for raw_row in rows[header_idx + 1:]:
        row = _repair_csv_row(raw_row, expected_cols)

        if len(row) <= met_col:
            continue
        cat = str(row[cat_col]).strip()
        met = str(row[met_col]).strip()
        if not cat or not met:
            continue
        key = (cat, met)

        row_sections = []
        for sec_start, sec_count in prev_sections:
            sec_vals = []
            for j in range(sec_start, sec_start + sec_count):
                val = str(row[j]).strip() if j < len(row) and row[j] else ""
                sec_vals.append(val)
            row_sections.append(sec_vals)
        hardcoded[key] = row_sections

        if op2_col is not None and op2_col < len(row):
            val = row[op2_col]
            if val and str(val).strip() not in ("", "NA", "-"):
                op2[key] = str(val).strip()
        if ytd_op2_col is not None and ytd_op2_col < len(row):
            val = row[ytd_op2_col]
            if val and str(val).strip() not in ("", "NA", "-"):
                ytd_op2[key] = str(val).strip()

    return op2, ytd_op2, hardcoded


# ═══════════════════════════════════════════════════════════════════
#  DATA STORE  — unified access to all loaded data
# ═══════════════════════════════════════════════════════════════════

class DataStore:
    """Holds all loaded data and provides a uniform get_value() interface."""

    def __init__(self):
        self.flat = {}
        self.multi = {}
        self.period_maps = {"weekly": {}, "monthly": {}, "pytd": {}}

    def add_flat(self, file_type, timeframe, data, period_map):
        self.flat[(file_type, timeframe)] = data
        self.period_maps[timeframe].update(period_map)

    def add_multi(self, file_type, timeframe, data, period_map):
        self.multi[(file_type, timeframe)] = data
        self.period_maps[timeframe].update(period_map)

    def get_value(self, device, file_type, sub_col, timeframe, label):
        """Get a single value for a device/metric/period."""
        key = (file_type, timeframe)

        if key in self.multi:
            dev_data = self.multi[key].get(device, {})
            period_data = dev_data.get(label, {})
            if isinstance(period_data, dict):
                return period_data.get(sub_col)
            return None

        if key in self.flat:
            dev_data = self.flat[key].get(device, {})
            return dev_data.get(label)

        return None

    def get_total(self, device, file_type, sub_col, timeframe):
        """Get the 'Total'/'YTD' value from a file."""
        return self.get_value(device, file_type, sub_col, timeframe, "__total__")


# ═══════════════════════════════════════════════════════════════════
#  FILE LOADING ORCHESTRATOR
# ═══════════════════════════════════════════════════════════════════

FILE_TYPE_CONFIG = {
    "dc":       {"format": "A"},
    "attach":   {"format": "A"},
    "reseller": {"format": "A"},
    "comp100":  {"format": "B", "sub_cols": ["Ovrl", "Base", "Swtnr"]},
    "comp75":   {"format": "B", "sub_cols": ["Ovrl", "Base"]},
    "depth":    {"format": "D"},
    "sweetner": {"format": "E"},
}

TIMEFRAMES = {
    "weekly": _parse_weekly,
    "monthly": _parse_monthly,
    "pytd": _parse_monthly,
}


def _load_all_files(input_dir, aliases_cfg):
    """Load all available files into a DataStore."""
    store = DataStore()
    input_dir = Path(input_dir)

    for ftype, cfg in FILE_TYPE_CONFIG.items():
        alias_map = _build_alias_map(aliases_cfg, ftype)
        fmt = cfg["format"]

        for tf, parse_fn in TIMEFRAMES.items():
            filepath = input_dir / f"{ftype}_{tf}.xlsx"
            if not filepath.is_file():
                continue

            print(f"  Loading {filepath.name}...")

            if fmt == "A":
                data, pm = _load_format_a(filepath, parse_fn, alias_map)
                store.add_flat(ftype, tf, data, pm)
            elif fmt == "B":
                data, pm = _load_format_b(filepath, parse_fn, alias_map, cfg["sub_cols"])
                store.add_multi(ftype, tf, data, pm)
            elif fmt == "D":
                data, pm = _load_format_d(filepath, parse_fn, alias_map)
                store.add_flat(ftype, tf, data, pm)
            elif fmt == "E":
                data, pm = _load_format_e(filepath, parse_fn, alias_map)
                store.add_multi(ftype, tf, data, pm)

    return store


# ═══════════════════════════════════════════════════════════════════
#  VALUE EXTRACTION HELPERS
# ═══════════════════════════════════════════════════════════════════

def _change_bps(v1, v2):
    """Basis point change: (v1 - v2) * 10000, rounded."""
    if v1 is not None and v2 is not None:
        try:
            return round((float(v1) - float(v2)) * 10000)
        except (ValueError, TypeError):
            pass
    return ""


def _fmt_pct(v):
    """Format a 0-1 decimal as a percentage string, e.g. 0.9571 -> '95.71%'."""
    if v is None:
        return ""
    try:
        return f"{float(v) * 100:.2f}%"
    except (ValueError, TypeError):
        return str(v)


def _parse_pct_str(s):
    """Parse a percentage string like '95.71%' or '7%' back to a float."""
    if not s or s in ("NA", "TBU", "-", "#VALUE!"):
        return None
    s = str(s).strip().replace(",", "").replace("%", "")
    try:
        return float(s) / 100.0
    except (ValueError, TypeError):
        return None


# ═══════════════════════════════════════════════════════════════════
#  ROW GENERATION
# ═══════════════════════════════════════════════════════════════════

def _build_rows(config, store, w_labels, m_labels, pytd_month_label,
                op2_vals, ytd_op2_vals, hardcoded_rows, total_cols, header):
    """Build all output rows from config-driven layout."""

    def pad(row):
        return row + [""] * (total_cols - len(row))

    row_num = 0

    def next_num():
        nonlocal row_num
        row_num += 1
        return row_num

    def get_raw(device, file_type, sub_col, timeframe, labels):
        """Get raw decimal values for a series of period labels."""
        return [store.get_value(device, file_type, sub_col, timeframe, l)
                for l in labels]

    def data_row(device, metric, file_type, sub_col):
        parts = ["", next_num(), device, metric]

        wv = get_raw(device, file_type, sub_col, "weekly", w_labels)
        wow = _change_bps(wv[0] if wv else None, wv[1] if len(wv) > 1 else None)
        parts += [wow] + [_fmt_pct(v) for v in wv]

        parts.append("")

        mv = get_raw(device, file_type, sub_col, "monthly", m_labels)
        mom = _change_bps(mv[0] if mv else None, mv[1] if len(mv) > 1 else None)
        parts += [mom] + [_fmt_pct(v) for v in mv]

        parts.append("")

        op2_key = (device, metric)
        op2_val = op2_vals.get(op2_key, "")
        op2_numeric = _parse_pct_str(op2_val) if op2_val else None
        latest_month_raw = mv[0] if mv else None
        vs_op2 = _change_bps(latest_month_raw, op2_numeric)
        parts += [op2_val, vs_op2]

        parts.append("")

        pytd_raw = store.get_value(device, file_type, sub_col, "pytd", pytd_month_label) if pytd_month_label else None
        mtd_yoy = _change_bps(latest_month_raw, pytd_raw)
        parts += [_fmt_pct(pytd_raw), mtd_yoy]

        ytd_raw = store.get_total(device, file_type, sub_col, "monthly")
        ytd_ly_raw = store.get_total(device, file_type, sub_col, "pytd")
        yoy = _change_bps(ytd_raw, ytd_ly_raw)
        parts += [_fmt_pct(ytd_raw), _fmt_pct(ytd_ly_raw), yoy]

        ytd_op2_val = ytd_op2_vals.get(op2_key, "")
        ytd_op2_numeric = _parse_pct_str(ytd_op2_val) if ytd_op2_val else None
        vs_ytd_op2 = _change_bps(ytd_raw, ytd_op2_numeric)
        parts += [ytd_op2_val, vs_ytd_op2]

        return pad(parts)

    def hardcoded_data_row(device, metric):
        """For 'Used item selection' — carry values from previous report.
        Uses section-based mapping so column count changes don't misalign data.
        """
        key = (device, metric)
        prev_secs = hardcoded_rows.get(key, [])

        parts = ["", next_num(), device, metric]
        cur_sections = _header_sections(header)
        pos = 4
        for si, (sec_start, sec_count) in enumerate(cur_sections):
            while pos < sec_start:
                parts.append("")
                pos += 1
            prev_vals = prev_secs[si] if si < len(prev_secs) else []
            for j in range(sec_count):
                parts.append(prev_vals[j] if j < len(prev_vals) else "")
                pos += 1

        return pad(parts)

    all_rows = []
    groups = config.get("groups", {})

    for group_name, group_cfg in groups.items():
        all_rows.append(pad([""] * total_cols))
        all_rows.append(pad(["", "", group_name]))

        for section in group_cfg.get("sections", []):
            sec_name = section["name"]
            all_rows.append(pad(["", "", "", sec_name]))

            for row_def in section.get("rows", []):
                device = row_def["device"]
                metric = row_def["metric"]
                file_type = row_def["file_type"]
                sub_col = row_def.get("sub_col")

                if file_type == "hardcoded":
                    all_rows.append(hardcoded_data_row(device, metric))
                else:
                    all_rows.append(data_row(device, metric, file_type, sub_col))

    return all_rows


# ═══════════════════════════════════════════════════════════════════
#  MAIN ENTRY POINT
# ═══════════════════════════════════════════════════════════════════

def run(config_path, input_dir, output_dir, **kwargs):
    """Entry point called by the dispatcher or standalone CLI."""
    with open(config_path) as f:
        config = yaml.safe_load(f)

    input_dir = Path(input_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n[Business Metrics] Loading files from {input_dir}...")

    aliases_cfg = config.get("device_aliases", {})
    store = _load_all_files(input_dir, aliases_cfg)

    output_path = output_dir / "report.csv"
    bootstrap_path = input_dir / "previous_report.csv"

    op2_vals, ytd_op2_vals, hardcoded_rows = {}, {}, {}
    for source in [output_path, bootstrap_path]:
        if not source.is_file():
            continue
        s_op2, s_ytd, s_hc = _load_previous_report(source)
        for k, v in s_hc.items():
            hardcoded_rows.setdefault(k, v)
        for k, v in s_op2.items():
            op2_vals.setdefault(k, v)
        for k, v in s_ytd.items():
            ytd_op2_vals.setdefault(k, v)
        print(f"  Previous report: {source.name} ({len(s_op2)} OP2, {len(s_hc)} rows)")

    w_pm = store.period_maps.get("weekly", {})
    m_pm = store.period_maps.get("monthly", {})
    p_pm = store.period_maps.get("pytd", {})

    w_labels = _sort_labels(w_pm)[:MAX_WEEKLY_PERIODS]
    m_labels = _sort_labels(m_pm)[:MAX_MONTHLY_PERIODS]

    today = date.today()
    cur_month = today.month

    pytd_month_label = None
    if p_pm:
        for label, (year, month) in p_pm.items():
            if month == cur_month:
                pytd_month_label = label
                break
        if pytd_month_label is None:
            pytd_month_label = _sort_labels(p_pm)[0] if p_pm else None

    current_year = today.year
    prev_year = current_year - 1

    latest_month_label = m_labels[0] if m_labels else ""
    op2_col_name = f"{latest_month_label} OP2" if latest_month_label else "OP2"

    header = ["", "#", "Category", "Metrics"]
    header += ["WoW"] + w_labels
    header += ["", "MoM"] + m_labels
    header += ["", op2_col_name, "v/s OP2"]
    pytd_label = f"{NUM_TO_MONTH.get(cur_month, 'MTD')}'{prev_year % 100:02d}"
    header += ["", pytd_label, "MTD YoY"]
    header += ["YTD", "YTD LY", "YoY"]
    header += ["YTD OP2", "v/s OP2"]
    total_cols = len(header)

    def pad(row):
        return row + [""] * (total_cols - len(row))

    rows = []
    title = config.get("title", "5. Business Metrics")
    rows.append(pad(["", "", title]))
    rows.append([""] * total_cols)
    rows.append(header)

    data_rows = _build_rows(
        config, store, w_labels, m_labels, pytd_month_label,
        op2_vals, ytd_op2_vals, hardcoded_rows, total_cols, header
    )
    rows.extend(data_rows)

    output_path = output_dir / "report.csv"
    with open(output_path, "w", newline="") as f:
        csv.writer(f).writerows(rows)

    print(f"\n[Business Metrics] Done!  {output_path}")
    print(f"  Weekly columns  : {len(w_labels)} ({', '.join(w_labels[:3])}...)")
    print(f"  Monthly columns : {len(m_labels)} ({', '.join(m_labels[:3])}...)")
    print(f"  PYTD month      : {pytd_month_label or 'none'}")
    print(f"  OP2 values      : {len(op2_vals)}")
    n_data = sum(1 for r in data_rows if r[1] and str(r[1]).strip().isdigit())
    print(f"  Data rows       : {n_data}")
    print(f"  Total columns   : {total_cols}")

    return str(output_path)


# ═══════════════════════════════════════════════════════════════════
#  STANDALONE CLI
# ═══════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    p = argparse.ArgumentParser(
        description="Business Metrics: Transform data into WBR report")
    args = p.parse_args()

    project_root = Path(__file__).resolve().parent
    run(
        config_path=project_root / "configs" / "business_metrics.yaml",
        input_dir=project_root / "input" / "business_metrics",
        output_dir=project_root / "output" / "business_metrics",
    )
