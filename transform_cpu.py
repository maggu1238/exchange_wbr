#!/usr/bin/env python3
"""
CPU metric: Transform raw CPU xlsx data into a WBR-style hierarchical report.

Can be run standalone:
    python3 transform_cpu.py
    python3 transform_cpu.py mobile tv --num-weeks 6

Or called by the dispatcher (run.py) via the run() entry point.
"""

import pandas as pd
import csv
import argparse
import re
import yaml
from pathlib import Path
from datetime import date

METRIC_NAME = "cpu"

# ═══════════════════════════════════════════════════════════════════
#  FIXED COLUMN MAPPINGS  (same CSV structure for every device)
# ═══════════════════════════════════════════════════════════════════

SEGMENT_DEFS = {
    "overall": {
        "label": None,
        "cpu_col": "CPU",
        "phases": [
            ("pre_pdd",  "Pre-PDD CPU",  "Pre-PDD"),
            ("pdd",      "PDD - CPU",    "PDD"),
            ("post_pdd", "Post-PDD CPU", "Post-PDD"),
        ],
    },
    "generalist": {
        "label": "Generalist",
        "cpu_col": "Gen",
        "phases": [
            ("pre_pdd",  "Pre-PDD CPU",  "Pre-PDD Gen"),
            ("pdd",      "PDD - CPU",    "PDD_Gen"),
            ("post_pdd", "Post-PDD CPU", "Post-PDD Gen"),
        ],
    },
    "specialist": {
        "label": "Specialist",
        "cpu_col": "Spl",
        "phases": [
            ("pre_pdd",  "Pre-PDD CPU",  "Pre-PDD Spl"),
            ("pdd",      "PDD - CPU",    "PDD Spl"),
            ("post_pdd", "Post-PDD CPU", "Post-PDD Spl"),
        ],
    },
}

MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_TO_NUM = {m: i for i, m in enumerate(MONTH_ABBR, 1)}
NUM_TO_MONTH = {i: m for m, i in MONTH_TO_NUM.items()}

TOP_N = 3

# ═══════════════════════════════════════════════════════════════════
#  CONFIG
# ═══════════════════════════════════════════════════════════════════

def _load_config(config_path):
    with open(config_path) as f:
        return yaml.safe_load(f)


def _active_segments(device_cfg):
    keys = device_cfg.get("segments", ["overall"])
    result = []
    for k in keys:
        if k not in SEGMENT_DEFS:
            raise SystemExit(f"Unknown segment '{k}'. "
                             f"Choose from: {', '.join(SEGMENT_DEFS)}")
        result.append((k, SEGMENT_DEFS[k]))
    return result


def _resolve_reasons(device_cfg, seg_key, phase_key, df, phase_col):
    cfg_reasons = (device_cfg
                   .get("reasons", {})
                   .get(seg_key, {})
                   .get(phase_key))
    if cfg_reasons is None or cfg_reasons == "auto":
        exclude = {"", "Total"}
        ranking = (
            df[~df["reason"].isin(exclude)]
            .groupby("reason")[phase_col]
            .sum()
            .sort_values(ascending=False)
        )
        return ranking.head(TOP_N).index.tolist()
    return list(cfg_reasons)


def _resolve_path(base_dir, path_str):
    p = Path(path_str)
    return str(p if p.is_absolute() else base_dir / p)


# ═══════════════════════════════════════════════════════════════════
#  AUTO-DISCOVERY from input/cpu/ folder
#  Pattern (case-insensitive): CPU_<device>_<timeframe>.xlsx
# ═══════════════════════════════════════════════════════════════════

DEVICE_ALIASES = {
    "mob": "mobile", "mobile": "mobile",
    "tv": "tv",
    "ac": "ac",
    "ref": "refrigerator", "refrigerator": "refrigerator",
    "wm": "wm",
}
DEVICE_FILE_TAGS = {
    "mobile": "mob", "tv": "tv", "ac": "ac",
    "refrigerator": "ref", "wm": "wm",
}
VALID_TIMEFRAMES = {"weekly", "monthly", "quarterly", "pytd"}
_DISCOVER_RE = re.compile(
    r"^cpu_([a-z]+)_([a-z]+)\.xlsx$", re.IGNORECASE
)


def _discover_input_files(input_dir):
    found = {}
    if not input_dir.is_dir():
        return found
    for f in input_dir.iterdir():
        m = _DISCOVER_RE.match(f.name)
        if not m:
            continue
        raw_device = m.group(1).lower()
        raw_tf = m.group(2).lower()
        device_key = DEVICE_ALIASES.get(raw_device)
        if device_key and raw_tf in VALID_TIMEFRAMES:
            found.setdefault(device_key, {})[raw_tf] = str(f.resolve())
    return found

# ═══════════════════════════════════════════════════════════════════
#  TIME-FRAME PARSERS  ->  ((year, period), display_label)
# ═══════════════════════════════════════════════════════════════════

def _parse_weekly(tf):
    parts = str(tf).split("-")
    year, week = int(parts[0]), int(parts[1])
    return (year, week), f"Wk-{week:02d}"


def _parse_monthly(tf):
    parts = str(tf).split("-")
    year, month = int(parts[0]), int(parts[1])
    label = f"{NUM_TO_MONTH[month]}'{year % 100:02d}"
    return (year, month), label


def _parse_quarterly(tf):
    parts = str(tf).split("-")
    year, q = int(parts[0]), int(parts[1])
    label = f"Q{q}'{year % 100:02d}"
    return (year, q), label


_parse_pytd = _parse_quarterly

# ═══════════════════════════════════════════════════════════════════
#  FILE LOADING  (xlsx wide-format: periods as column groups)
# ═══════════════════════════════════════════════════════════════════

METRICS = ["CPU", "Pre-PDD", "PDD", "Post-PDD", "Gen", "Pre-PDD Gen",
           "PDD_Gen", "Post-PDD Gen", "Spl", "Pre-PDD Spl", "PDD Spl",
           "Post-PDD Spl"]
BLOCK_SIZE = len(METRICS)


def _load_file(filepath, parse_fn):
    import openpyxl
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=True)
    header0 = list(next(rows_iter))
    next(rows_iter)

    periods = []
    for i, val in enumerate(header0):
        if val is not None and i > 0 and str(val).strip().lower() != "total":
            periods.append((str(val).strip(), i))

    records = []
    for row in rows_iter:
        row = list(row)
        reason = str(row[0]).strip() if row[0] is not None else ""
        for period_label, col_start in periods:
            vals = row[col_start:col_start + BLOCK_SIZE]
            rec = {"reason": reason, "Time_Frame": period_label}
            for metric, val in zip(METRICS, vals):
                rec[metric] = val
            records.append(rec)

    wb.close()

    df = pd.DataFrame(records)
    df["reason"] = df["reason"].fillna("").str.strip()
    for col in METRICS:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    parsed = df["Time_Frame"].apply(parse_fn)
    df["_sort_key"] = parsed.apply(lambda x: x[0])
    df["_label"] = parsed.apply(lambda x: x[1])
    return df


def _sorted_labels(df):
    pairs = df.drop_duplicates("_label")[["_label", "_sort_key"]]
    return pairs.sort_values("_sort_key", ascending=False)["_label"].tolist()


def _drop_future_periods(df, granularity):
    today = date.today()
    cur_year, cur_month = today.year, today.month

    if granularity == "monthly":
        cutoff = (cur_year, cur_month)
    elif granularity in ("quarterly", "pytd"):
        cur_quarter = (cur_month - 1) // 3 + 1
        cutoff = (cur_year, cur_quarter)
    else:
        return df

    mask = df["_sort_key"].apply(lambda k: k <= cutoff)
    dropped = (~mask).sum()
    if dropped:
        labels = df.loc[~mask, "_label"].unique().tolist()
        print(f"  [filter] Dropped {dropped} future rows from {granularity}: {labels}")
    return df[mask].copy()


def _load_device_files(device_key, device_cfg, input_dir,
                       cli_overrides=None):
    disc = _discover_input_files(input_dir).get(device_key, {})
    ov = cli_overrides or {}
    files_cfg = device_cfg.get("files", {})
    required_tfs = set(device_cfg.get("timeframes", ["weekly"]))

    def pick(key):
        return ov.get(key) or disc.get(key) or files_cfg.get(key)

    missing = []
    for tf in sorted(required_tfs):
        path = pick(tf)
        resolved = _resolve_path(input_dir, path) if path else None
        if not resolved or not Path(resolved).is_file():
            tag = DEVICE_FILE_TAGS.get(device_key, device_key)
            missing.append(f"  - {tf:12s}  (expected: CPU_{tag}_{tf}.xlsx)")

    if missing:
        raise SystemExit(
            f"\nDevice '{device_key}': missing required input file(s) "
            f"in input/cpu/ folder:\n" + "\n".join(missing) + "\n"
        )

    all_parsers = {
        "weekly": _parse_weekly,
        "monthly": _parse_monthly,
        "quarterly": _parse_quarterly,
        "pytd": _parse_pytd,
    }

    result = {}
    for tf, parser in all_parsers.items():
        path = pick(tf)
        resolved = _resolve_path(input_dir, path) if path else None
        if resolved and Path(resolved).is_file():
            df = _load_file(resolved, parser)
            result[tf] = _drop_future_periods(df, tf) if tf != "weekly" else df
        else:
            result[tf] = None

    return result

# ═══════════════════════════════════════════════════════════════════
#  VALUE EXTRACTION
# ═══════════════════════════════════════════════════════════════════

def _change_bps(vals):
    if len(vals) >= 2 and vals[0] != "" and vals[1] != "":
        try:
            return round((float(vals[0]) - float(vals[1])) * 10000)
        except (ValueError, TypeError):
            pass
    return "-"


AGG_ROW_NAME = "Total"


def _agg_values(df, col, labels):
    total_df = df[df["reason"] == AGG_ROW_NAME]
    if not total_df.empty:
        agg = total_df.set_index("_label")[col]
        return [agg.get(l, 0) if pd.notna(agg.get(l)) else "" for l in labels]
    agg = df.groupby("_label")[col].sum()
    return [agg.get(l, 0) for l in labels]


def _reason_values(df, reason, col, labels):
    rdf = df[df["reason"] == reason].set_index("_label")[col]
    out = []
    for l in labels:
        v = rdf.get(l, None)
        out.append(v if pd.notna(v) else "")
    return out


def _compute_avg(values, labels, year_filter=None):
    selected = []
    for v, l in zip(values, labels):
        if v == "":
            continue
        if year_filter is not None:
            m = re.match(r"Q(\d)'(\d{2,4})", l)
            if m:
                yr = int(m.group(2)) + (2000 if len(m.group(2)) == 2 else 0)
                if yr != year_filter:
                    continue
        selected.append(float(v))
    if selected:
        return sum(selected) / len(selected)
    return ""

# ═══════════════════════════════════════════════════════════════════
#  PER-DEVICE ROW GENERATION
# ═══════════════════════════════════════════════════════════════════

def _generate_device_rows(device_cfg, dfs, w_labels, m_labels, q_labels,
                          p_labels, current_year, total_cols):
    weekly_df = dfs["weekly"]
    monthly_df = dfs.get("monthly")
    quarterly_df = dfs.get("quarterly")
    pytd_df = dfs.get("pytd")

    segments = _active_segments(device_cfg)
    title = device_cfg.get("title", "CPU")

    def pad(row):
        return row + [""] * (total_cols - len(row))

    row_num = 0

    def next_num():
        nonlocal row_num
        row_num += 1
        return f"{row_num:.1f}"

    def _vals(df, col, labels, is_reason, reason):
        if df is None:
            return [""] * len(labels)
        if is_reason:
            return _reason_values(df, reason, col, labels)
        return _agg_values(df, col, labels)

    def data_row(label, col, is_reason=False, reason=None):
        parts = ["", next_num(), label]

        wv = _vals(weekly_df, col, w_labels, is_reason, reason)
        parts += [_change_bps(wv)] + wv

        if m_labels:
            mv = _vals(monthly_df, col, m_labels, is_reason, reason)
            parts += ["", _change_bps(mv)] + mv

        if q_labels:
            qv = _vals(quarterly_df, col, q_labels, is_reason, reason)
            parts += ["", _change_bps(qv)] + qv

        if q_labels or pytd_df is not None:
            parts += [""]
            if q_labels:
                parts += [_compute_avg(qv, q_labels, year_filter=current_year)]
            if pytd_df is not None:
                pv = _vals(pytd_df, col, p_labels, is_reason, reason)
                parts += [_compute_avg(pv, p_labels)]

        return pad(parts)

    rows = []
    rows.append(pad(["", "", title]))

    for seg_key, seg_def in segments:
        if seg_def["label"]:
            rows.append(pad(["", next_num(), seg_def["label"]]))

        rows.append(data_row("Contacts per unit (CPU)", seg_def["cpu_col"]))

        for phase_key, phase_label, phase_col in seg_def["phases"]:
            rows.append(data_row(phase_label, phase_col))

            reasons = _resolve_reasons(
                device_cfg, seg_key, phase_key, weekly_df, phase_col
            )
            for rank, reason in enumerate(reasons, 1):
                rows.append(data_row(
                    f"Reason {rank} -{reason}", phase_col,
                    is_reason=True, reason=reason,
                ))

    return rows

# ═══════════════════════════════════════════════════════════════════
#  MAIN ENTRY POINT
# ═══════════════════════════════════════════════════════════════════

def run(config_path, input_dir, output_dir, devices=None, num_weeks=None,
        cli_overrides=None):
    """Entry point called by the dispatcher or standalone CLI.

    Args:
        config_path: Path to configs/cpu.yaml
        input_dir:   Path to input/cpu/
        output_dir:  Path to output/cpu/
        devices:     List of device keys to process (None = all from config)
        num_weeks:   Limit number of week columns shown
        cli_overrides: Dict of {timeframe: filepath} overrides
    """
    full_config = _load_config(config_path)
    all_devices = full_config.get("devices", {})

    input_dir = Path(input_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    device_keys = devices if devices else list(all_devices.keys())

    for dk in device_keys:
        if dk not in all_devices:
            raise SystemExit(
                f"Device '{dk}' not in config. "
                f"Available: {', '.join(all_devices)}"
            )

    single = len(device_keys) == 1

    # ── Load data for every device ──
    device_data = {}
    for dk in device_keys:
        dcfg = all_devices[dk]
        ov = cli_overrides if single else None
        dfs = _load_device_files(dk, dcfg, input_dir, ov)
        device_data[dk] = dfs

    # ── Build column layout from first device ──
    first_dfs = device_data[device_keys[0]]

    w_labels = _sorted_labels(first_dfs["weekly"])
    if num_weeks:
        w_labels = w_labels[:num_weeks]

    m_labels = _sorted_labels(first_dfs["monthly"]) if first_dfs["monthly"] is not None else []
    q_labels = _sorted_labels(first_dfs["quarterly"]) if first_dfs["quarterly"] is not None else []
    p_labels = _sorted_labels(first_dfs["pytd"]) if first_dfs["pytd"] is not None else []

    current_year = max(k[0] for k in first_dfs["weekly"]["_sort_key"])
    prev_year = current_year - 1

    # ── Header ──
    header = ["", "#", "Metrics"]
    header += ["WoW"] + w_labels
    if m_labels:
        header += ["", "MoM"] + m_labels
    if q_labels:
        header += ["", "QoQ"] + q_labels
    has_pytd = first_dfs["pytd"] is not None
    if q_labels or has_pytd:
        header += [""]
        if q_labels:
            header += [f"YTD'{current_year % 100:02d}"]
        if has_pytd:
            header += [f"PYTD'{prev_year % 100:02d}"]
    total_cols = len(header)

    def pad(row):
        return row + [""] * (total_cols - len(row))

    # ── Assemble output ──
    rows = []
    rows.append(pad(["", "", "Exchange WBR - CPU"]))
    rows.append([""] * total_cols)
    rows.append(header)

    for dk in device_keys:
        dcfg = all_devices[dk]
        dfs = device_data[dk]

        if num_weeks:
            dfs["weekly"] = dfs["weekly"][
                dfs["weekly"]["_label"].isin(w_labels)
            ].copy()

        rows.append([""] * total_cols)
        device_rows = _generate_device_rows(
            dcfg, dfs, w_labels, m_labels, q_labels, p_labels,
            current_year, total_cols,
        )
        rows.extend(device_rows)

    # ── Write CSV ──
    output_path = output_dir / "report.csv"
    with open(output_path, "w", newline="") as f:
        csv.writer(f).writerows(rows)

    # ── Summary ──
    print(f"\n[CPU] Done!  {output_path}")
    print(f"  Devices  : {', '.join(device_keys)}")
    for dk in device_keys:
        dcfg = all_devices[dk]
        segs = [s[0] for s in _active_segments(dcfg)]
        manual = sum(
            1 for sk, sd in _active_segments(dcfg)
            for pk, _, _ in sd["phases"]
            if isinstance(dcfg.get("reasons", {}).get(sk, {}).get(pk), list)
        )
        auto = sum(len(sd["phases"]) for _, sd in _active_segments(dcfg)) - manual
        print(f"    {dk:12s} : segments=[{', '.join(segs)}]  "
              f"reasons: {manual} config / {auto} auto")
    sections = [f"weekly ({len(w_labels)})"]
    if m_labels:
        sections.append(f"monthly ({len(m_labels)})")
    if q_labels:
        sections.append(f"quarterly ({len(q_labels)} + YTD Avg)")
    if p_labels:
        sections.append(f"PYTD ({len(p_labels)} qtrs -> avg)")
    print(f"  Timeline : {' | '.join(sections)}")
    print(f"  Total    : {sum(1 for r in rows if r[1] and r[1] != '#')} "
          f"metric rows | {total_cols} columns")

    return str(output_path)


# ═══════════════════════════════════════════════════════════════════
#  STANDALONE CLI
# ═══════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    p = argparse.ArgumentParser(
        description="CPU metric: Transform data into WBR report")

    p.add_argument("devices", nargs="*",
                   help="Device key(s) from config (e.g. mobile tv). "
                        "Omit to process all devices.")

    p.add_argument("--weekly", help="Weekly xlsx override (single-device only)")
    p.add_argument("--monthly", help="Monthly xlsx override (single-device only)")
    p.add_argument("--quarterly", help="Quarterly xlsx override (single-device only)")
    p.add_argument("--pytd", help="PYTD xlsx override (single-device only)")
    p.add_argument("--num-weeks", type=int,
                   help="Limit number of week columns shown")

    args = p.parse_args()

    project_root = Path(__file__).resolve().parent

    cli_overrides = {}
    if args.weekly:
        cli_overrides["weekly"] = args.weekly
    if args.monthly:
        cli_overrides["monthly"] = args.monthly
    if args.quarterly:
        cli_overrides["quarterly"] = args.quarterly
    if args.pytd:
        cli_overrides["pytd"] = args.pytd

    run(
        config_path=project_root / "configs" / "cpu.yaml",
        input_dir=project_root / "input" / "cpu",
        output_dir=project_root / "output" / "cpu",
        devices=args.devices or None,
        num_weeks=args.num_weeks,
        cli_overrides=cli_overrides or None,
    )
